"""Real documents, built by hand for the tests and the benchmarks: a Word file, a spreadsheet, a presentation and a PDF,
each the smallest that the readers markitdown uses (mammoth, openpyxl, python-pptx, pdfminer) accept as genuine."""

import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

_CONTENT_TYPES = ('<?xml version="1.0" encoding="UTF-8"?>'
                  '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                  '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                  '<Default Extension="xml" ContentType="application/xml"/>'
                  '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument'
                  '.wordprocessingml.document.main+xml"/></Types>')
_RELS = ('<?xml version="1.0" encoding="UTF-8"?>'
         '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
         '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
         'officeDocument" Target="word/document.xml"/></Relationships>')


def make_docx(path: Path, paragraphs: list[tuple[str, str]]) -> Path:
    """A Word document of (style, text) paragraphs; a style is "Heading1", "Heading2" or "Normal"."""
    body = "".join(f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr><w:r><w:t xml:space="preserve">{escape(text)}'
                   "</w:t></w:r></w:p>" for style, text in paragraphs)
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                f"{body}</w:body></w:document>")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", _CONTENT_TYPES)
        package.writestr("_rels/.rels", _RELS)
        package.writestr("word/document.xml", document)
    return path


def make_xlsx(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    """A workbook: each sheet's name and its rows (the first row is the header)."""
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(path)
    return path


def make_pptx(path: Path, slides: list[tuple[str, str]]) -> Path:
    """A presentation of (title, text) slides."""
    from pptx import Presentation

    deck = Presentation()
    for title, text in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])  # a title and a text box
        slide.shapes.title.text = title
        slide.placeholders[1].text = text
    deck.save(path)
    return path


def _pdf_text(line: str) -> str:
    """A line as a PDF string's contents."""
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages: list[list[str]]) -> Path:
    """A PDF with one line of Helvetica text per string, a page per list (Latin-1 text only)."""
    count = len(pages)
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(count))}] /Count {count} >>".encode(),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, lines in enumerate(pages):
        content = ("BT /F1 11 Tf 14 TL 56 780 Td " + " ".join(f"({_pdf_text(line)}) Tj T*" for line in lines)
                   + " ET").encode("latin-1", "replace")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
                       f"/Contents {5 + 2 * i} 0 R >>".encode())
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    table = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{table}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path
